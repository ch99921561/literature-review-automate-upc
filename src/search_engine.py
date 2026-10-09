# -*- coding: utf-8 -*-
"""
search_engine.py - Motor de búsqueda para Literature Review
"""

import os
import itertools
import time
from collections import deque
from datetime import datetime
from typing import Dict, List, Optional, Tuple, Any
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .config import (
    APIType, API_CONFIGS, OUTPUTS_DIR, CONSOLIDATED_OUTPUT_PREFIX
)
from .models import (
    SearchFilters, ScopusFilters, IEEEFilters, WOSFilters,
    SearchResult, CombinationResult
)
from .logger import logger
from .base_client import BaseAPIClient
from .input_config import InputConfig


class RequestLimiter:
    """Controla cuotas opcionales de solicitudes durante una ejecución."""

    def __init__(self, calls_per_second: Optional[int],
                 calls_per_day: Optional[int]) -> None:
        self.calls_per_second = calls_per_second
        self.calls_per_day = calls_per_day
        self.calls_made = 0
        self._recent_calls = deque()

    def acquire(self) -> bool:
        """Espera si es necesario y reserva una solicitud, si queda cuota."""
        if self.calls_per_day is not None and self.calls_made >= self.calls_per_day:
            return False

        if self.calls_per_second is not None:
            now = time.monotonic()
            while self._recent_calls and now - self._recent_calls[0] >= 1:
                self._recent_calls.popleft()
            if len(self._recent_calls) >= self.calls_per_second:
                time.sleep(1 - (now - self._recent_calls[0]))
                return self.acquire()
            self._recent_calls.append(time.monotonic())

        self.calls_made += 1
        return True


class SearchEngine:
    """Motor de búsqueda que coordina múltiples clientes de APIs."""
    
    def __init__(self):
        self.clients: Dict[APIType, BaseAPIClient] = {}
        self.config: Optional[InputConfig] = None
        self.individual_results: Dict[APIType, List[SearchResult]] = {}
        self.unexecuted_combinations: Dict[APIType, List[Tuple[str, ...]]] = {}
        self.zero_result_combinations: Dict[APIType, List[CombinationResult]] = {}
    
    def register_client(self, api_type: APIType, client: BaseAPIClient) -> bool:
        """Registra un cliente de API si está autenticado correctamente."""
        if client.authenticate():
            self.clients[api_type] = client
            print(f"  ✓ {client.get_api_name()} autenticado correctamente")
            return True
        else:
            print(f"  ✗ {client.get_api_name()} no autenticado (API_KEY no configurada)")
            return False
    
    def load_config(self, thesis_id: Optional[str] = None) -> bool:
        """Carga la configuración desde el archivo de entrada."""
        self.config = InputConfig.load(thesis_id=thesis_id)
        return self.config is not None
    
    def run_simple_mode(self, api_type: APIType) -> Tuple[int, List[CombinationResult]]:
        """
        Ejecuta el modo sencillo (conteo de publicaciones).
        
        Returns:
            Tupla (código_retorno, lista_combinaciones)
        """
        if api_type not in self.clients:
            logger.write(f"ERROR: Cliente {api_type.value} no registrado")
            return (1, [])
        
        client = self.clients[api_type]
        api_name = client.get_api_name()
        
        # Inicializar log
        log_filename = logger.init(api_name, "sencilla")
        logger.write(f"Archivo de log: {log_filename}")
        
        logger.header(f"MODO SENCILLO - {api_name.upper()}")
        
        # Obtener filtros según el tipo de API
        filters = self._get_filters_for_api(api_type)
        limiter = self._get_request_limiter(api_type)
        
        # Mostrar configuración
        self._print_config(api_type, filters)
        
        # Ejecutar búsquedas
        individual_results = self._search_individual(client, filters, limiter)
        self.individual_results[api_type] = individual_results
        combination_results, unexecuted_combinations = self._search_combinations(
            client, filters, limiter, individual_results
        )
        self.unexecuted_combinations[api_type] = unexecuted_combinations
        self.zero_result_combinations[api_type] = self._zero_result_combinations(
            combination_results
        )
        
        # Guardar resultados
        self._save_results(api_type, individual_results, combination_results)
        
        logger.close()
        return (0, combination_results)
    
    def _get_filters_for_api(self, api_type: APIType) -> SearchFilters:
        """Obtiene los filtros específicos para una API."""
        if api_type == APIType.SCOPUS:
            return self.config.scopus
        elif api_type == APIType.IEEE:
            return self.config.ieee
        elif api_type == APIType.WOS:
            return self.config.wos
        return SearchFilters(year_from=self.config.year_from, year_to=self.config.year_to)

    def _get_request_limiter(self, api_type: APIType) -> RequestLimiter:
        """Crea un limitador a partir de la configuración opcional de cuotas."""
        limits = self.config.rate_limits.get(api_type.value, {})
        return RequestLimiter(
            calls_per_second=limits.get("calls_per_second"),
            calls_per_day=limits.get("calls_per_day"),
        )
    
    def _print_config(self, api_type: APIType, filters: SearchFilters) -> None:
        """Imprime la configuración cargada."""
        logger.header("CONFIGURACIÓN CARGADA")
        logger.write(f"Archivo: definitions/input.json")
        if self.config.thesis_id:
            logger.write(f"Tesis seleccionada: {self.config.thesis_id}")
            logger.write(f"Descripción: {self.config.thesis_description}")
        logger.write(f"Keywords: {len(self.config.keywords)}")
        for i, kw in enumerate(self.config.keywords, 1):
            logger.write(f"  {i}. {kw}")
        
        logger.write(f"\nFiltros:")
        logger.write(f"  Años: {filters.year_from or 'Sin límite'} - {filters.year_to or 'Sin límite'}")
        
        if isinstance(filters, ScopusFilters):
            logger.write(f"  Tipos de documento: {', '.join(filters.doc_types) if filters.doc_types else 'Todos'}")
            logger.write(f"  Áreas temáticas: {', '.join(filters.subject_areas) if filters.subject_areas else 'Todas'}")
        elif isinstance(filters, IEEEFilters):
            logger.write(f"  Tipos de contenido: {', '.join(filters.content_types) if filters.content_types else 'Todos'}")
            limits = self.config.rate_limits.get(APIType.IEEE.value, {})
            logger.write(f"  Límite de solicitudes/segundo: {limits.get('calls_per_second') or 'Sin límite'}")
            logger.write(f"  Límite de solicitudes/día: {limits.get('calls_per_day') or 'Sin límite'}")
        elif isinstance(filters, WOSFilters):
            logger.write(f"  Base de datos: {filters.database}")
            logger.write(f"  Edición: {filters.edition or 'Todas'}")
            logger.write(f"  Tipos de documento: {', '.join(filters.document_types) if filters.document_types else 'Todos'}")
    
    def _search_individual(self, client: BaseAPIClient, 
                           filters: SearchFilters,
                           limiter: RequestLimiter) -> List[SearchResult]:
        """Realiza búsqueda individual por keyword."""
        logger.header("RESULTADOS INDIVIDUALES")
        logger.write(f"{'Keyword':<50} | {'Publicaciones':>15}")
        logger.write("-" * 70)
        
        results = []
        total = 0
        
        for keyword in self.config.keywords:
            if not limiter.acquire():
                logger.write("Límite diario alcanzado; se omiten las keywords restantes.")
                break
            query = f'"{keyword}"'
            count = client.count_results(query, filters)
            
            if count == -1:
                logger.write(f"{keyword:<50} | {'ERROR':>15}")
                results.append(SearchResult(keyword=keyword, query=query, count=None, error=True))
            else:
                logger.write(f"{keyword:<50} | {count:>15,}")
                results.append(SearchResult(keyword=keyword, query=query, count=count))
                total += count
            
            time.sleep(0.25)
        
        logger.write("-" * 70)
        logger.write(f"{'TOTAL INDIVIDUAL (suma)':<50} | {total:>15,}")
        
        return results
    
    def _search_combinations(self, client: BaseAPIClient,
                              filters: SearchFilters,
                              limiter: RequestLimiter,
                              individual_results: List[SearchResult]) -> Tuple[List[CombinationResult], List[Tuple[str, ...]]]:
        """Realiza búsqueda por combinaciones de 3 keywords."""
        keywords = self.config.keywords
        
        if len(keywords) < 3:
            logger.write("\nNOTA: Se necesitan al menos 3 keywords para generar combinaciones.")
            return [], []
        
        logger.header("COMBINACIONES DE 3 KEYWORDS (TERNAS)")
        
        unique_keywords = self._unique_keywords(keywords)
        all_combinations = list(itertools.combinations(unique_keywords, 3))
        combinations = self._eligible_combinations(keywords, individual_results)
        skipped_combinations = len(all_combinations) - len(combinations)

        logger.write(f"Total de combinaciones posibles: {len(all_combinations)}")
        logger.write(
            "Combinaciones omitidas (al menos una keyword con 0 resultados): "
            f"{skipped_combinations}"
        )
        logger.write(f"Combinaciones a consultar: {len(combinations)}")
        logger.write("")
        
        results = []
        unexecuted_combinations = []
        total = 0
        
        for idx, combo in enumerate(combinations, 1):
            if not limiter.acquire():
                logger.write(
                    f"\nLímite diario alcanzado después de {limiter.calls_made} solicitudes; "
                    "se omiten las combinaciones restantes."
                )
                unexecuted_combinations = combinations[idx - 1:]
                break
            query = f'"{combo[0]}" AND "{combo[1]}" AND "{combo[2]}"'
            count = client.count_results(query, filters)
            
            display_keywords = f"[{combo[0]}] AND [{combo[1]}] AND [{combo[2]}]"
            
            if count == -1:
                logger.write(f"\n{idx:3}. ERROR")
                logger.write(f"     Keywords: {display_keywords}")
                logger.write(f"     Query enviada: {query}")
                results.append(CombinationResult(keywords=list(combo), query=query, count=None, error=True))
            else:
                logger.write(f"\n{idx:3}. Resultados: {count:,}")
                logger.write(f"     Keywords: {display_keywords}")
                logger.write(f"     Query enviada: {query}")
                results.append(CombinationResult(keywords=list(combo), query=query, count=count))
                total += count
            
            time.sleep(0.25)

        if unexecuted_combinations:
            individual_counts = {
                result.keyword: result.count
                for result in individual_results
            }
            logger.header("TERNAS ELEGIBLES NO EJECUTADAS")
            logger.write(
                "No se ejecutaron por alcanzar el límite de solicitudes de la API:"
            )
            for combo in unexecuted_combinations:
                combination_with_counts = " AND ".join(
                    f"{keyword} ({individual_counts[keyword]})"
                    for keyword in combo
                )
                logger.write(f"  {combination_with_counts}")

        zero_result_combinations = self._zero_result_combinations(results)
        if zero_result_combinations:
            individual_counts = {
                result.keyword: result.count
                for result in individual_results
            }
            logger.header("TERNAS SIN RESULTADOS")
            logger.write(
                "Las siguientes ternas tenían resultados individuales positivos, "
                "pero devolvieron 0 al aplicar AND:"
            )
            for combination in zero_result_combinations:
                combination_with_counts = " AND ".join(
                    f"{keyword} ({individual_counts[keyword]})"
                    for keyword in combination.keywords
                )
                logger.write(f"  {combination_with_counts}")
        
        # Mostrar resumen y TOP 30
        self._print_combination_summary(results, total, client, filters, limiter)
        
        return results, unexecuted_combinations

    @staticmethod
    def _unique_keywords(keywords: List[str]) -> List[str]:
        """Elimina keywords duplicadas sin distinguir mayúsculas y minúsculas."""
        unique_keywords = []
        seen_keywords = set()
        for keyword in keywords:
            normalized_keyword = keyword.casefold()
            if normalized_keyword not in seen_keywords:
                unique_keywords.append(keyword)
                seen_keywords.add(normalized_keyword)
        return unique_keywords

    @classmethod
    def _eligible_combinations(cls, keywords: List[str],
                               individual_results: List[SearchResult]) -> List[Tuple[str, ...]]:
        """Retorna ternas cuyas tres keywords tienen conteo individual positivo."""
        keywords_with_results = {
            result.keyword
            for result in individual_results
            if result.count is not None and result.count > 0
        }
        return [
            combination
            for combination in itertools.combinations(cls._unique_keywords(keywords), 3)
            if all(keyword in keywords_with_results for keyword in combination)
        ]
    
    def _print_combination_summary(self, results: List[CombinationResult], total: int,
                                   client: BaseAPIClient = None, 
                                   filters: SearchFilters = None,
                                   limiter: RequestLimiter = None) -> None:
        """Imprime el resumen de combinaciones."""
        logger.header("RESUMEN DE COMBINACIONES")
        logger.write(f"Total de combinaciones: {len(results)}")
        logger.write(f"Suma de resultados: {total:,}")
        
        # Filtrar combinaciones con resultados > 0
        with_results = [r for r in results if r.count and r.count > 0]
        logger.write(f"Combinaciones con al menos 1 resultado: {len(with_results)}")
        
        if with_results:
            with_results.sort(key=lambda x: x.count or 0, reverse=True)
            
            # Obtener documentos para el TOP 30 si tenemos cliente y filtros
            top_30 = with_results[:30]
            if client and filters:
                logger.write("")
                logger.write("Obteniendo títulos de documentos para el TOP 30...")
                for idx, r in enumerate(top_30, 1):
                    if limiter and not limiter.acquire():
                        logger.write(
                            "Límite diario alcanzado; no se recuperarán más documentos."
                        )
                        break
                    documents = client.get_documents(r.query, filters, max_docs=200)
                    r.documents = documents
                    logger.write(f"  Llave {idx}: {len(documents)} documentos obtenidos")
                    time.sleep(0.25)
            
            logger.header("TOP 30 COMBINACIONES CON MÁS RESULTADOS")
            logger.write("")
            logger.write(f"{'Llave':<6} | {'Resultados':>12} | {'Keyword 1':<25} | {'Keyword 2':<25} | {'Keyword 3':<25}")
            logger.write("-" * 102)
            
            for i, r in enumerate(top_30, 1):
                k1 = r.keywords[0][:24] if len(r.keywords[0]) > 24 else r.keywords[0]
                k2 = r.keywords[1][:24] if len(r.keywords[1]) > 24 else r.keywords[1]
                k3 = r.keywords[2][:24] if len(r.keywords[2]) > 24 else r.keywords[2]
                logger.write(f"{i:<6} | {r.count:>12,} | {k1:<25} | {k2:<25} | {k3:<25}")
            
            logger.write("-" * 102)
            logger.write("")
            logger.write("Detalle de queries enviadas:")
            for i, r in enumerate(top_30, 1):
                logger.write(f"  {i:2}. {r.query}")
            
            # Tabla: DOCUMENTOS POR LLAVE
            if any(r.documents for r in top_30):
                logger.header("DOCUMENTOS ENCONTRADOS POR LLAVE (TOP 30)")
                logger.write("")
                for i, r in enumerate(top_30, 1):
                    logger.write(f"{'='*80}")
                    logger.write(f"LLAVE {i} - {len(r.documents)} documento(s)")
                    logger.write(f"Keywords: {' AND '.join(r.keywords)}")
                    logger.write(f"{'='*80}")
                    if r.documents:
                        for doc_idx, document in enumerate(r.documents, 1):
                            title = document["titulo"]
                            display_title = title[:120] + "..." if len(title) > 120 else title
                            year = document["año_publicacion"] or "No disponible"
                            logger.write(f"  {doc_idx:3}. {display_title} ({year})")
                    else:
                        logger.write("  (Sin documentos recuperados)")
                    logger.write("")
    
    def _build_documents_by_key(self, combinations: List[CombinationResult]) -> List[Dict[str, Any]]:
        """Construye la tabla de documentos por llave (TOP 30)."""
        with_results = [r for r in combinations if r.count and r.count > 0]
        with_results.sort(key=lambda x: x.count or 0, reverse=True)
        
        documents_table = []
        for i, r in enumerate(with_results[:30], 1):
            documents_table.append({
                "llave": i,
                "keywords": r.keywords,
                "query": r.query,
                "count": r.count,
                "documentos": r.documents
            })
        
        return documents_table

    @staticmethod
    def _unused_keywords(keywords: List[str],
                         individual_results: List[SearchResult],
                         combinations: List[CombinationResult]) -> List[Dict[str, Any]]:
        """Identifica keywords que no participaron en ninguna terna ejecutada."""
        individual_counts = {
            result.keyword.casefold(): result.count
            for result in individual_results
        }
        used_keywords = {
            keyword.casefold()
            for combination in combinations
            for keyword in combination.keywords
        }
        unused = []
        for keyword in SearchEngine._unique_keywords(keywords):
            normalized_keyword = keyword.casefold()
            if normalized_keyword in used_keywords:
                continue

            count = individual_counts.get(normalized_keyword)
            if count is None:
                reason = "No se obtuvo el conteo individual"
            elif count == 0:
                reason = "0 resultados individuales"
            else:
                reason = "No participó en una terna ejecutada"
            unused.append({
                "keyword": keyword,
                "individual_count": count,
                "reason": reason,
            })
        return unused

    @staticmethod
    def _zero_result_combinations(
        combinations: List[CombinationResult],
    ) -> List[CombinationResult]:
        """Retorna ternas ejecutadas que devolvieron cero resultados."""
        return [
            combination
            for combination in combinations
            if combination.count == 0 and not combination.error
        ]
    
    def _save_results(self, api_type: APIType, 
                      individual: List[SearchResult],
                      combinations: List[CombinationResult]) -> None:
        """Guarda los resultados en archivo JSON."""
        config = API_CONFIGS[api_type]
        filters = self._get_filters_for_api(api_type)
        
        # Calcular totales
        total_individual = sum(r.count or 0 for r in individual)
        total_combinations = sum(r.count or 0 for r in combinations)
        
        output_data = {
            "api": api_type.value,
            "mode": "sencilla",
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "input_file": "definitions/input.json",
            "thesis": {
                "id": self.config.thesis_id,
                "description": self.config.thesis_description,
            },
            "filters": {
                "year_from": filters.year_from,
                "year_to": filters.year_to,
            },
            "individual_results": {
                "keywords": [
                    {"keyword": r.keyword, "query": r.query, "count": r.count, "error": r.error}
                    for r in individual
                ],
                "total": total_individual,
            },
            "combination_results": {
                "combination_size": 3,
                "total_combinations": len(combinations),
                "combinations": [
                    {"keywords": r.keywords, "query": r.query, "count": r.count, "error": r.error, "documents": r.documents}
                    for r in combinations
                ],
                "total": total_combinations,
            },
            "documents_by_key": self._build_documents_by_key(combinations),
        }
        
        # Agregar filtros específicos
        if isinstance(filters, ScopusFilters):
            output_data["filters"]["doc_types"] = filters.doc_types
            output_data["filters"]["subject_areas"] = filters.subject_areas
        elif isinstance(filters, IEEEFilters):
            output_data["filters"]["content_types"] = filters.content_types
        
        # Guardar en carpeta outputs (la ruta ya incluye OUTPUTS_DIR)
        with open(config.output_counts_file, "w", encoding="utf-8") as f:
            json.dump(output_data, f, indent=2, ensure_ascii=False)
        
        logger.header("RESUMEN FINAL")
        logger.write(f"Keywords analizados: {len(self.config.keywords)}")
        logger.write(f"Total individual: {total_individual:,}")
        if combinations:
            logger.write(f"Combinaciones (ternas): {len(combinations)}")
            logger.write(f"Total combinaciones: {total_combinations:,}")
        logger.write(f"\nResultados guardados en: {config.output_counts_file}")
        logger.write(f"Log guardado en: {logger.filename}")
    
    def save_consolidated_top30(self, all_results: Dict[APIType, List[CombinationResult]]) -> str:
        """
        Guarda un archivo consolidado con el TOP 30 en formato Excel (.xlsx).
        
        Args:
            all_results: Diccionario {APIType: List[CombinationResult]} con resultados por API
        
        Returns:
            Nombre del archivo generado
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.join(OUTPUTS_DIR, f"{CONSOLIDATED_OUTPUT_PREFIX}_{timestamp}.xlsx")
        
        # Determinar qué APIs fueron ejecutadas
        apis_executed = list(all_results.keys())
        apis_names = [api.value for api in apis_executed]
        
        # Crear workbook
        wb = Workbook()
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Hoja de resumen
        ws_summary = wb.active
        ws_summary.title = "Resumen"
        ws_summary['A1'] = "TOP 30 COMBINACIONES - REPORTE CONSOLIDADO"
        ws_summary['A1'].font = Font(bold=True, size=14)
        ws_summary['A3'] = "Fecha y hora:"
        ws_summary['B3'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ws_summary['A4'] = "APIs ejecutadas:"
        ws_summary['B4'] = ', '.join(apis_names)
        ws_summary['A5'] = "Keywords:"
        ws_summary['B5'] = len(self.config.keywords)
        ws_summary['A6'] = "Rango de años:"
        ws_summary['B6'] = f"{self.config.year_from or 'Sin límite'} - {self.config.year_to or 'Sin límite'}"
        ws_summary['A7'] = "Tesis seleccionada:"
        ws_summary['B7'] = self.config.thesis_id or "No especificada"
        ws_summary['A8'] = "Descripción de tesis:"
        ws_summary['B8'] = self.config.thesis_description or "No especificada"

        ws_unused = wb.create_sheet(title="Keywords_Sin_Combinacion")
        unused_headers = ["API", "Keyword", "Resultados individuales", "Motivo"]
        for col, header in enumerate(unused_headers, 1):
            cell = ws_unused.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        unused_row = 2
        for api_type, combinations in all_results.items():
            unused_keywords = self._unused_keywords(
                self.config.keywords,
                self.individual_results.get(api_type, []),
                combinations,
            )
            for unused_keyword in unused_keywords:
                ws_unused.cell(row=unused_row, column=1, value=api_type.value).border = border
                ws_unused.cell(row=unused_row, column=2, value=unused_keyword["keyword"]).border = border
                ws_unused.cell(
                    row=unused_row,
                    column=3,
                    value=unused_keyword["individual_count"],
                ).border = border
                ws_unused.cell(row=unused_row, column=4, value=unused_keyword["reason"]).border = border
                unused_row += 1

        ws_unused.column_dimensions['A'].width = 14
        ws_unused.column_dimensions['B'].width = 45
        ws_unused.column_dimensions['C'].width = 25
        ws_unused.column_dimensions['D'].width = 40

        ws_unexecuted = wb.create_sheet(title="Ternas_No_Ejecutadas")
        unexecuted_headers = [
            "API",
            "Keyword 1",
            "Resultados individuales 1",
            "Keyword 2",
            "Resultados individuales 2",
            "Keyword 3",
            "Resultados individuales 3",
            "Motivo",
        ]
        for col, header in enumerate(unexecuted_headers, 1):
            cell = ws_unexecuted.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        unexecuted_row = 2
        for api_type, combinations in self.unexecuted_combinations.items():
            individual_counts = {
                result.keyword.casefold(): result.count
                for result in self.individual_results.get(api_type, [])
            }
            for combination in combinations:
                ws_unexecuted.cell(row=unexecuted_row, column=1, value=api_type.value).border = border
                for keyword_index, keyword in enumerate(combination):
                    column = 2 + (keyword_index * 2)
                    ws_unexecuted.cell(row=unexecuted_row, column=column, value=keyword).border = border
                    ws_unexecuted.cell(
                        row=unexecuted_row,
                        column=column + 1,
                        value=individual_counts.get(keyword.casefold()),
                    ).border = border
                ws_unexecuted.cell(
                    row=unexecuted_row,
                    column=8,
                    value="Límite de solicitudes alcanzado",
                ).border = border
                unexecuted_row += 1

        ws_unexecuted.column_dimensions["A"].width = 14
        for column in "BDF":
            ws_unexecuted.column_dimensions[column].width = 35
        for column in "CEG":
            ws_unexecuted.column_dimensions[column].width = 24
        ws_unexecuted.column_dimensions["H"].width = 38

        ws_zero_results = wb.create_sheet(title="Ternas_Sin_Resultados")
        zero_headers = [
            "API",
            "Keyword 1",
            "Resultados individuales 1",
            "Keyword 2",
            "Resultados individuales 2",
            "Keyword 3",
            "Resultados individuales 3",
            "Resultado de la terna",
        ]
        for col, header in enumerate(zero_headers, 1):
            cell = ws_zero_results.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center')

        zero_row = 2
        for api_type, combinations in self.zero_result_combinations.items():
            individual_counts = {
                result.keyword.casefold(): result.count
                for result in self.individual_results.get(api_type, [])
            }
            for combination in combinations:
                ws_zero_results.cell(row=zero_row, column=1, value=api_type.value).border = border
                for keyword_index, keyword in enumerate(combination.keywords):
                    column = 2 + (keyword_index * 2)
                    ws_zero_results.cell(row=zero_row, column=column, value=keyword).border = border
                    ws_zero_results.cell(
                        row=zero_row,
                        column=column + 1,
                        value=individual_counts.get(keyword.casefold()),
                    ).border = border
                ws_zero_results.cell(row=zero_row, column=8, value=0).border = border
                zero_row += 1

        ws_zero_results.column_dimensions["A"].width = 14
        for column in "BDF":
            ws_zero_results.column_dimensions[column].width = 35
        for column in "CEG":
            ws_zero_results.column_dimensions[column].width = 24
        ws_zero_results.column_dimensions["H"].width = 24
        
        # Crear hojas por API
        for api_type, combinations in all_results.items():
            api_name = api_type.value.upper()
            
            # Filtrar combinaciones con resultados > 0 y ordenar
            with_results = [r for r in combinations if r.count and r.count > 0]
            with_results.sort(key=lambda x: x.count or 0, reverse=True)
            
            if not with_results:
                continue
            
            # Hoja de combinaciones TOP 30
            ws_combo = wb.create_sheet(title=f"{api_name}_TOP30")
            
            # Headers
            headers = ["Rank", "Resultados", "Keyword 1", "Keyword 2", "Keyword 3", "Query"]
            for col, header in enumerate(headers, 1):
                cell = ws_combo.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
            
            # Datos
            for i, r in enumerate(with_results[:30], 1):
                row = i + 1
                ws_combo.cell(row=row, column=1, value=i).border = border
                ws_combo.cell(row=row, column=2, value=r.count).border = border
                ws_combo.cell(row=row, column=3, value=r.keywords[0]).border = border
                ws_combo.cell(row=row, column=4, value=r.keywords[1]).border = border
                ws_combo.cell(row=row, column=5, value=r.keywords[2]).border = border
                ws_combo.cell(row=row, column=6, value=r.query).border = border
            
            # Ajustar anchos
            ws_combo.column_dimensions['A'].width = 8
            ws_combo.column_dimensions['B'].width = 12
            ws_combo.column_dimensions['C'].width = 30
            ws_combo.column_dimensions['D'].width = 30
            ws_combo.column_dimensions['E'].width = 30
            ws_combo.column_dimensions['F'].width = 80
            
            # Hoja de documentos
            ws_docs = wb.create_sheet(title=f"{api_name}_Documentos")
            
            # Headers documentos
            doc_headers = [
                "Llave",
                "Keywords",
                "Titulo",
                "Año de publicación",
                "Posible URL de descarga",
                "API_Source",
            ]
            for col, header in enumerate(doc_headers, 1):
                cell = ws_docs.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
                cell.alignment = Alignment(horizontal='center')
            
            # Datos documentos
            doc_row = 2
            for i, r in enumerate(with_results[:30], 1):
                if r.documents:
                    keywords_str = " AND ".join(r.keywords)
                    for document in r.documents:
                        ws_docs.cell(row=doc_row, column=1, value=i).border = border
                        ws_docs.cell(row=doc_row, column=2, value=keywords_str).border = border
                        ws_docs.cell(row=doc_row, column=3, value=document["titulo"]).border = border
                        ws_docs.cell(row=doc_row, column=4, value=document["año_publicacion"] or "No disponible").border = border
                        ws_docs.cell(
                            row=doc_row,
                            column=5,
                            value=document.get("url_descarga") or "No disponible",
                        ).border = border
                        ws_docs.cell(row=doc_row, column=6, value=api_type.value).border = border
                        doc_row += 1
            
            # Ajustar anchos documentos
            ws_docs.column_dimensions['A'].width = 8
            ws_docs.column_dimensions['B'].width = 60
            ws_docs.column_dimensions['C'].width = 100
            ws_docs.column_dimensions['D'].width = 20
            ws_docs.column_dimensions['E'].width = 80
            ws_docs.column_dimensions['F'].width = 12
        
        # Guardar archivo
        wb.save(filename)
        
        print(f"\n{'='*80}")
        print(f"  ARCHIVO CONSOLIDADO GENERADO (XLSX)")
        print(f"{'='*80}")
        print(f"Archivo: {filename}")
        print(f"APIs incluidas: {', '.join(apis_names)}")
        
        # Mostrar resumen
        for api_type, combinations in all_results.items():
            with_results = [r for r in combinations if r.count and r.count > 0]
            with_results.sort(key=lambda x: x.count or 0, reverse=True)
            
            if with_results:
                print(f"\n[{api_type.value.upper()}] TOP 5 (de {len(with_results)} con resultados):")
                for i, r in enumerate(with_results[:5], 1):
                    keywords_str = " AND ".join(r.keywords)
                    print(f"  {i:2}. {r.count:,} resultados - {keywords_str}")
            else:
                print(f"\n[{api_type.value.upper()}] Sin combinaciones con resultados")
            
            unexecuted_count = len(self.unexecuted_combinations.get(api_type, []))
            if unexecuted_count:
                print(
                    f"[{api_type.value.upper()}] {unexecuted_count} terna(s) elegible(s) "
                    "no ejecutada(s); revisa la hoja Ternas_No_Ejecutadas."
                )
            
            zero_result_count = len(self.zero_result_combinations.get(api_type, []))
            if zero_result_count:
                print(
                    f"[{api_type.value.upper()}] {zero_result_count} terna(s) con "
                    "resultados individuales positivos devolvieron 0; revisa la "
                    "hoja Ternas_Sin_Resultados."
                )
        
        return filename


def run_extended_mode(engine: SearchEngine) -> int:
    """Ejecuta el modo extendido interactivo."""
    print("\n--- Selecciona la API ---")
    print("1. Scopus")
    print("2. IEEE Xplore")
    print("3. Web of Science")
    
    api_choice = input("\nSelecciona API (1, 2 o 3): ").strip()
    
    if api_choice == "1":
        api_type = APIType.SCOPUS
    elif api_choice == "2":
        api_type = APIType.IEEE
    elif api_choice == "3":
        api_type = APIType.WOS
    else:
        print("Opción no válida")
        return 1
    
    if api_type not in engine.clients:
        print(f"ERROR: Cliente {api_type.value} no disponible")
        return 1
    
    client = engine.clients[api_type]
    config = API_CONFIGS[api_type]
    
    logger.header(f"MODO EXTENDIDO - {api_type.value.upper()}")
    
    query = input("\nIngresa tu búsqueda (ej: 'machine learning AND healthcare'): ").strip()
    if not query:
        query = "machine learning AND systematic review"
        print(f"Usando query por defecto: {query}")
    
    # Filtros básicos
    print("\n--- Filtros opcionales (Enter para omitir) ---")
    year_from_str = input("Año desde (ej: 2020): ").strip()
    year_from = int(year_from_str) if year_from_str.isdigit() else None
    
    year_to_str = input("Año hasta (ej: 2025): ").strip()
    year_to = int(year_to_str) if year_to_str.isdigit() else None
    
    # Crear filtros según API
    if api_type == APIType.SCOPUS:
        filters = ScopusFilters(year_from=year_from, year_to=year_to)
    elif api_type == APIType.IEEE:
        filters = IEEEFilters(year_from=year_from, year_to=year_to)
    elif api_type == APIType.WOS:
        filters = WOSFilters(year_from=year_from, year_to=year_to)
    else:
        filters = SearchFilters(year_from=year_from, year_to=year_to)
    
    # Modo de búsqueda
    print("\n--- Modo de búsqueda ---")
    print(f"1. Búsqueda simple (hasta {config.max_per_request} resultados)")
    print("2. Obtener TODOS los resultados (con paginación)")
    mode = input("Selecciona modo (1 o 2, default 1): ").strip()
    
    if mode == "2":
        max_str = input("Máximo de resultados (default 200): ").strip()
        max_results = int(max_str) if max_str.isdigit() else 200
        
        all_entries = client.search_all(query, filters, max_results)
        
        logger.header("RESUMEN DE RESULTADOS")
        print(f"Total obtenido: {len(all_entries)} artículos")
        
        if all_entries:
            print("\nPrimeros 5 resultados:")
            for i, entry in enumerate(all_entries[:5], 1):
                title = entry.get('dc:title') or entry.get('title', 'N/A')
                print(f"  {i}. {title[:80]}...")
        
        output_data = {
            "api": api_type.value,
            "query": query,
            "filters": {"year_from": year_from, "year_to": year_to},
            "total_results": len(all_entries),
            "entries": all_entries,
        }
        
        with open(config.output_results_file, "w", encoding="utf-8") as f:
            json.dump(output_data, f, indent=2, ensure_ascii=False)
        print(f"\nResultados guardados en: {config.output_results_file}")
    else:
        count_str = input(f"Número de resultados (máx {config.max_per_request}, default 25): ").strip()
        max_records = int(count_str) if count_str.isdigit() else 25
        
        response = client.search(query, filters, max_records, verbose=True)
        
        # Mostrar resultados
        total = client.parse_total_results(response)
        entries = client.parse_entries(response)
        
        logger.header("RESULTADOS DE BÚSQUEDA")
        print(f"Total de resultados: {total}")
        print(f"Mostrando: {len(entries)}\n")
        
        for i, entry in enumerate(entries, 1):
            title = entry.get('dc:title') or entry.get('title', 'N/A')
            print(f"--- Resultado {i} ---")
            print(f"Título: {title}")
            print()
        
        with open(config.output_results_file, "w", encoding="utf-8") as f:
            json.dump(response, f, indent=2, ensure_ascii=False)
        print(f"\nRespuesta guardada en: {config.output_results_file}")
    
    return 0
